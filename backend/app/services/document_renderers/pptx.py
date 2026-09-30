"""Deterministic 16:9 PPTX rendering from a validated content plan."""

from __future__ import annotations

from io import BytesIO
from typing import Any, Literal

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from app.services.document_generation_content import PURPOSE_NOTICE, DocumentContentPlan

_BLUE = RGBColor(31, 111, 179)
_DARK = RGBColor(35, 55, 75)
_MUTED = RGBColor(96, 115, 133)
_BUILTIN_COVER = RGBColor(25, 78, 112)
_BUILTIN_BODY = RGBColor(238, 247, 249)
_BUILTIN_LIGHT = RGBColor(226, 241, 246)

PptxTemplateMode = Literal["default", "builtin"]


def _refs(source_ids: list[int]) -> str:
    return " ".join(f"[来源{item}]" for item in dict.fromkeys(source_ids))


def _all_source_ids(plan: DocumentContentPlan) -> list[int]:
    return list(dict.fromkeys(item for section in plan.sections for item in section.source_ids))


def _set_run(run, size: int, color: RGBColor = _DARK, *, bold: bool = False) -> None:
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold


def _footer(slide, refs: str = "", *, color: RGBColor = _MUTED) -> None:
    box = slide.shapes.add_textbox(Inches(0.45), Inches(7.08), Inches(12.4), Inches(0.22))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = f"AI 生成的健康科普资料｜不作诊断依据  {refs}".strip()
    _set_run(run, 8, color)


def _title(slide, text: str, *, color: RGBColor = _BLUE) -> None:
    box = slide.shapes.add_textbox(Inches(0.65), Inches(0.35), Inches(12), Inches(0.7))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = text
    _set_run(run, 26, color, bold=True)


def _background(slide, color: RGBColor) -> None:
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def render_pptx(
    plan: DocumentContentPlan,
    sources: list[dict[str, Any]],
    *,
    template_mode: PptxTemplateMode = "default",
    template_bytes: bytes | None = None,
) -> tuple[bytes, int]:
    if template_mode not in {"default", "builtin"}:
        raise ValueError("Unsupported PPT template mode")
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]

    cover = presentation.slides.add_slide(blank)
    if template_mode == "builtin":
        _background(cover, _BUILTIN_COVER)
    title = cover.shapes.add_textbox(Inches(1), Inches(1.6), Inches(11.3), Inches(1.4))
    paragraph = title.text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = plan.title
    _set_run(run, 34, _BUILTIN_LIGHT if template_mode == "builtin" else _BLUE, bold=True)
    subtitle = cover.shapes.add_textbox(Inches(1.5), Inches(3.2), Inches(10.3), Inches(1.3))
    notices = [f"适用读者：{plan.audience}", PURPOSE_NOTICE]
    if any(source.get("trustLevel") == "prototype" for source in sources):
        notices.append("开发原型资料，未经医学专业审核")
    for index, value in enumerate(notices):
        item = (
            subtitle.text_frame.paragraphs[0] if index == 0 else subtitle.text_frame.add_paragraph()
        )
        item.alignment = PP_ALIGN.CENTER
        run = item.add_run()
        run.text = value
        _set_run(
            run,
            16 if index == 0 else 12,
            _BUILTIN_LIGHT if template_mode == "builtin" else _MUTED,
            bold=index == 1,
        )
    _footer(cover, color=_BUILTIN_LIGHT if template_mode == "builtin" else _MUTED)

    summary = presentation.slides.add_slide(blank)
    if template_mode == "builtin":
        _background(summary, _BUILTIN_BODY)
    _title(summary, "内容摘要")
    body = summary.shapes.add_textbox(Inches(0.9), Inches(1.45), Inches(11.5), Inches(4.9))
    body.text_frame.word_wrap = True
    run = body.text_frame.paragraphs[0].add_run()
    run.text = plan.summary
    _set_run(run, 22)
    _footer(summary, _refs(_all_source_ids(plan)))

    for section in plan.sections:
        slide = presentation.slides.add_slide(blank)
        if template_mode == "builtin":
            _background(slide, _BUILTIN_BODY)
        _title(slide, section.heading)
        body = slide.shapes.add_textbox(Inches(0.85), Inches(1.3), Inches(11.7), Inches(5.45))
        frame = body.text_frame
        frame.word_wrap = True
        values = [*section.paragraphs, *section.bullets]
        for index, value in enumerate(values[:6]):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.level = 0
            paragraph.space_after = Pt(10)
            run = paragraph.add_run()
            prefix = "• " if value in section.bullets else ""
            run.text = f"{prefix}{value}"
            _set_run(run, 18 if len(values) <= 4 else 16)
        _footer(slide, _refs(section.source_ids))

    if plan.warning_signs:
        warning = presentation.slides.add_slide(blank)
        if template_mode == "builtin":
            _background(warning, _BUILTIN_BODY)
        _title(warning, "需要及时就医的情况")
        body = warning.shapes.add_textbox(Inches(0.85), Inches(1.3), Inches(11.7), Inches(5.45))
        frame = body.text_frame
        frame.word_wrap = True
        refs: list[int] = []
        for index, item in enumerate(plan.warning_signs):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            run = paragraph.add_run()
            run.text = f"• {item.text}"
            _set_run(run, 17)
            refs.extend(item.source_ids)
        _footer(warning, _refs(refs))

    references = presentation.slides.add_slide(blank)
    if template_mode == "builtin":
        _background(references, _BUILTIN_BODY)
    _title(references, "参考资料")
    body = references.shapes.add_textbox(Inches(0.65), Inches(1.15), Inches(12), Inches(5.8))
    frame = body.text_frame
    frame.word_wrap = True
    for index, source in enumerate(sources):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        run = paragraph.add_run()
        run.text = (
            f"[来源{source['referenceNumber']}] {source['organization']}｜"
            f"{source['title']}\n{source['canonicalUrl']}"
        )
        _set_run(run, 12)
    _footer(references)

    target = BytesIO()
    presentation.save(target)
    output = target.getvalue()
    if template_bytes is not None:
        from app.services.presentation_template import apply_presentation_template

        output = apply_presentation_template(output, template_bytes)
    return output, len(presentation.slides)
