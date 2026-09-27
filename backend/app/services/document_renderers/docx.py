"""Deterministic DOCX rendering from a validated content plan."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from app.services.document_generation_content import PURPOSE_NOTICE, DocumentContentPlan


def _refs(source_ids: list[int]) -> str:
    return " ".join(f"[来源{item}]" for item in dict.fromkeys(source_ids))


def _all_source_ids(plan: DocumentContentPlan) -> list[int]:
    return list(
        dict.fromkeys(item for section in plan.sections for item in section.source_ids)
    )


def _font(run, size: int = 11, *, bold: bool = False) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.bold = bold


def render_docx(plan: DocumentContentPlan, sources: list[dict[str, Any]]) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    properties = document.core_properties
    properties.author = "病理诊断科普助手"
    properties.title = "AI 医学健康科普资料"
    properties.subject = "健康科普"
    properties.comments = PURPOSE_NOTICE

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(title.add_run(plan.title), 22, bold=True)
    audience = document.add_paragraph()
    audience.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(audience.add_run(f"适用读者：{plan.audience}"), 10)
    notice = document.add_paragraph()
    notice.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(notice.add_run(PURPOSE_NOTICE), 9, bold=True)
    if any(source.get("trustLevel") == "prototype" for source in sources):
        prototype = document.add_paragraph()
        prototype.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _font(prototype.add_run("开发原型资料，未经医学专业审核"), 9, bold=True)

    document.add_heading("内容摘要", level=1)
    paragraph = document.add_paragraph()
    _font(paragraph.add_run(f"{plan.summary} {_refs(_all_source_ids(plan))}"), 11)

    for item in plan.sections:
        document.add_heading(item.heading, level=1)
        refs = _refs(item.source_ids)
        for value in item.paragraphs:
            paragraph = document.add_paragraph()
            _font(paragraph.add_run(f"{value} {refs}"), 11)
        for value in item.bullets:
            paragraph = document.add_paragraph(style="List Bullet")
            _font(paragraph.add_run(f"{value} {refs}"), 11)

    if plan.warning_signs:
        document.add_heading("需要及时就医的情况", level=1)
        for item in plan.warning_signs:
            paragraph = document.add_paragraph(style="List Bullet")
            _font(paragraph.add_run(f"{item.text} {_refs(item.source_ids)}"), 11)

    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_heading("参考资料", level=1)
    for source in sources:
        paragraph = document.add_paragraph()
        text = (
            f"[来源{source['referenceNumber']}] {source['organization']}：《{source['title']}》\n"
            f"{source['canonicalUrl']}"
        )
        _font(paragraph.add_run(text), 9)

    for section in document.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _font(footer.add_run("AI 生成的健康科普资料｜不作诊断依据"), 8)

    target = BytesIO()
    document.save(target)
    return target.getvalue()
