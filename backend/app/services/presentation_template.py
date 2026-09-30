"""Validate and reuse a simple PPTX as visual-only cover/body slide designs.

Template bytes never enter the content model. Editable text on reused slides is
cleared; slide masters, backgrounds, pictures, drawing geometry and page size
remain in the output.
"""

from __future__ import annotations

from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile

from lxml import etree
from pptx import Presentation
from pptx.enum.dml import MSO_COLOR_TYPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt

from app.core.errors import AppError

_REL_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_MAX_BYTES = 10 * 1024 * 1024
_MAX_UNPACKED = 40 * 1024 * 1024
_MAX_PARTS = 600


def validate_presentation_template(content: bytes) -> None:
    if not content or len(content) > _MAX_BYTES:
        raise AppError(413, "PPT_TEMPLATE_TOO_LARGE", "PPT 模板不能超过 10 MB")
    try:
        with ZipFile(BytesIO(content)) as archive:
            parts = archive.infolist()
            if len(parts) > _MAX_PARTS or not parts:
                raise ValueError("too many parts")
            if len({part.filename.lower() for part in parts}) != len(parts):
                raise ValueError("duplicate parts")
            total = 0
            for part in parts:
                if part.flag_bits & 1:
                    raise ValueError("encrypted archive")
                name = part.filename.replace("\\", "/").lower()
                if name.startswith("/") or ".." in name.split("/"):
                    raise ValueError("unsafe path")
                if any(
                    x in name
                    for x in ("vbaproject", "activex/", "embeddings/", "oleobject", "ppt/media/")
                ):
                    # Pictures are in ppt/media and are allowed; audio/video is checked below.
                    if not name.startswith("ppt/media/") or not name.endswith(
                        (".png", ".jpg", ".jpeg", ".webp", ".gif")
                    ):
                        raise ValueError("active or unsupported content")
                if name.endswith((".svg", ".emf", ".wmf", ".mp4", ".mov", ".avi")):
                    raise ValueError("unsupported media")
                total += part.file_size
                if total > _MAX_UNPACKED or (
                    part.compress_size and part.file_size / part.compress_size > 100
                ):
                    raise ValueError("archive bomb")
                if name.endswith(".rels"):
                    rels = etree.fromstring(
                        archive.read(part),
                        parser=etree.XMLParser(resolve_entities=False, no_network=True),
                    )
                    if any(rel.get("TargetMode", "").lower() == "external" for rel in rels):
                        raise ValueError("external relationship")
        deck = Presentation(BytesIO(content))
        if len(deck.slides) < 2 or len(deck.slides) > 20:
            raise ValueError("cover and body slide required")
        for slide in deck.slides:
            for shape in slide.shapes:
                if shape.has_chart or shape.has_table:
                    raise ValueError("unsupported shape")
                _check_image_relationships(shape.element, slide.part)
            if slide._element.cSld.bg is not None:
                _check_image_relationships(slide._element.cSld.bg, slide.part)
    except Exception as exc:
        raise AppError(
            422,
            "PPT_TEMPLATE_INVALID",
            "请选择含封面和正文示例页的普通 PPTX 模板；不支持外链、宏、音视频或嵌入对象",
        ) from exc


def _check_image_relationships(element, part) -> None:
    for node in element.iter():
        for attr, relationship_id in node.attrib.items():
            if attr.startswith(_REL_NS):
                if not part.rels[relationship_id].reltype.endswith("/image"):
                    raise ValueError("unsupported slide relationship")


def _copy_design_element(element, source_part, target_part):
    copied = deepcopy(element)
    for node in copied.iter():
        for attr, old_id in list(node.attrib.items()):
            if not attr.startswith(_REL_NS):
                continue
            rel = source_part.rels[old_id]
            if not rel.reltype.endswith("/image"):
                raise ValueError("unsupported slide relationship")
            node.set(attr, target_part.relate_to(rel.target_part, rel.reltype))
    return copied


def _clone_slide(deck, donor):
    slide = deck.slides.add_slide(donor.slide_layout)
    for shape in list(slide.shapes):
        slide.shapes._spTree.remove(shape.element)
    donor_bg = donor._element.cSld.bg
    if donor_bg is not None:
        target_bg = slide._element.cSld.bg
        if target_bg is not None:
            slide._element.cSld.remove(target_bg)
        slide._element.cSld.insert(0, _copy_design_element(donor_bg, donor.part, slide.part))
    for shape in donor.shapes:
        element = _copy_design_element(shape.element, donor.part, slide.part)
        slide.shapes._spTree.insert_element_before(element, "p:extLst")
    return slide


def _replace_content(slide, source, *, cover: bool, slide_width: int, slide_height: int) -> None:
    slots = [
        shape for shape in slide.shapes if shape.has_text_frame and shape.width > slide_width // 4
    ]
    slots.sort(key=lambda shape: shape.top)
    texts = [shape.text for shape in source.shapes if shape.has_text_frame and shape.text.strip()]
    slot_colors = []
    for shape in slots:
        color = None
        for paragraph in shape.text_frame.paragraphs:
            if paragraph.runs:
                font_color = paragraph.runs[0].font.color
                if font_color.type == MSO_COLOR_TYPE.RGB:
                    color = font_color.rgb
                break
        slot_colors.append(color)
    for shape in slide.shapes:
        if shape.has_text_frame:
            shape.text_frame.clear()
    if not texts:
        return
    target_slots = slots[:2]
    for index, text in enumerate(texts[:2]):
        if index < len(target_slots):
            target = target_slots[index]
            frame = target.text_frame
        else:
            left = int(slide_width * 0.07)
            top = int(slide_height * (0.2 if index == 0 else 0.36))
            target = slide.shapes.add_textbox(
                left,
                top,
                int(slide_width * 0.86),
                int(slide_height * (0.16 if index == 0 else 0.48)),
            )
            frame = target.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        frame.paragraphs[0].text = text
        for paragraph in frame.paragraphs:
            for run in paragraph.runs:
                run.font.name = "Microsoft YaHei"
                run.font.size = Pt(30 if cover and index == 0 else 22 if index == 0 else 17)
                if index < len(slot_colors) and slot_colors[index] is not None:
                    run.font.color.rgb = slot_colors[index]
    # The purpose notice and reference markers must remain visible regardless of template design.
    footer = texts[-1] if len(texts) > 2 else "AI 生成的健康科普资料｜不作诊断依据"
    box = slide.shapes.add_textbox(
        int(slide_width * 0.04), int(slide_height * 0.93), int(slide_width * 0.92), Inches(0.3)
    )
    box.text = footer
    for run in box.text_frame.paragraphs[0].runs:
        run.font.name = "Microsoft YaHei"
        run.font.size = Pt(9)


def apply_presentation_template(generated: bytes, template: bytes) -> bytes:
    validate_presentation_template(template)
    original = Presentation(BytesIO(generated))
    deck = Presentation(BytesIO(template))
    donors = list(deck.slides)
    source_count = len(deck.slides)
    for index, source in enumerate(original.slides):
        target = _clone_slide(deck, donors[min(index, source_count - 1)])
        _replace_content(
            target,
            source,
            cover=index == 0,
            slide_width=deck.slide_width,
            slide_height=deck.slide_height,
        )
    for _ in range(source_count):
        slide_id = deck.slides._sldIdLst[0]
        deck.part.drop_rel(slide_id.rId)
        deck.slides._sldIdLst.remove(slide_id)
    output = BytesIO()
    deck.save(output)
    return output.getvalue()
