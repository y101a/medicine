"""Isolated document parser worker.

The parent process sends bounded bytes over stdin and kills this process on timeout.
Only JSON metadata and extracted text are written to stdout; input bytes are never written to disk.
"""

import argparse
import io
import json
import math
import os
import re
import shutil
import subprocess
import sys
import zipfile
from dataclasses import asdict, dataclass
from pathlib import Path

from docx import Document
from PIL import Image, UnidentifiedImageError
from pptx import Presentation
from pypdf import PdfReader

MAX_ZIP_ENTRIES = 2_000
MAX_ZIP_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_ZIP_MEMBER_BYTES = 20 * 1024 * 1024
MAX_ZIP_COMPRESSION_RATIO = 100
MAX_OCR_IMAGES = 12
MAX_OCR_IMAGES_PER_PAGE = 3
MAX_OCR_IMAGE_BYTES = 8 * 1024 * 1024
MAX_OCR_IMAGE_PIXELS = 12_000_000
MAX_OCR_TEXT_CHARS = 4_000
OCR_IMAGE_TIMEOUT_SECONDS = 3


@dataclass(frozen=True)
class Section:
    locator: str
    text: str


@dataclass(frozen=True)
class TableCell:
    row: int
    column: int
    text: str


@dataclass(frozen=True)
class ExtractedTable:
    locator: str
    row_count: int
    column_count: int
    layout_uncertain: bool
    cells: tuple[TableCell, ...]


def table_sections(table: ExtractedTable) -> list[Section]:
    sections = []
    for row_index in range(1, table.row_count + 1):
        cells = [cell for cell in table.cells if cell.row == row_index]
        if any(cell.text for cell in cells):
            detail = "；".join(f"第 {cell.column} 列：{cell.text or '（空）'}" for cell in cells)
            warning = "（合并单元格，行列对应需核对）" if table.layout_uncertain else ""
            sections.append(
                Section(
                    f"{table.locator} 第 {row_index} 行",
                    f"{table.locator} 第 {row_index} 行{warning}：{detail}",
                )
            )
    return sections


class ParseFailure(Exception):
    def __init__(self, code: str, message: str) -> None:
        self.code = code
        self.message = message
        super().__init__(message)


def fail(code: str, message: str) -> None:
    raise ParseFailure(code, message)


def normalize_text(value: str) -> str:
    value = value.replace("\x00", " ")
    value = re.sub(r"[\u0001-\u0008\u000b\u000c\u000e-\u001f\u007f]", " ", value)
    value = re.sub(r"[ \t\u3000]+", " ", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def inspect_ooxml(content: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            infos = archive.infolist()
            if not infos or len(infos) > MAX_ZIP_ENTRIES:
                fail("FILE_ZIP_UNSAFE", "文档压缩结构不安全")
            total = 0
            names: set[str] = set()
            for info in infos:
                name = info.filename.replace("\\", "/")
                path_parts = [part for part in name.split("/") if part]
                if (
                    name.startswith("/")
                    or ".." in path_parts
                    or name.casefold().endswith("vbaproject.bin")
                    or name in names
                    or info.flag_bits & 0x1
                ):
                    fail("FILE_ZIP_UNSAFE", "文档压缩结构不安全")
                names.add(name)
                total += info.file_size
                if info.file_size > MAX_ZIP_MEMBER_BYTES or total > MAX_ZIP_UNCOMPRESSED_BYTES:
                    fail("FILE_ZIP_UNSAFE", "文档解压后内容过大")
                if info.compress_size == 0:
                    if info.file_size > 0:
                        fail("FILE_ZIP_UNSAFE", "文档压缩比异常")
                elif info.file_size / info.compress_size > MAX_ZIP_COMPRESSION_RATIO:
                    fail("FILE_ZIP_UNSAFE", "文档压缩比异常")
            if "[Content_Types].xml" not in names:
                fail("FILE_TYPE_NOT_ALLOWED", "仅支持 PDF、DOCX、PPTX、TXT 和 Markdown")
            if "word/document.xml" in names:
                return "docx"
            if "ppt/presentation.xml" in names:
                return "pptx"
    except ParseFailure:
        raise
    except (zipfile.BadZipFile, OSError, ValueError) as exc:
        raise ParseFailure("FILE_PARSE_FAILED", "文档已损坏或无法解析") from exc
    fail("FILE_TYPE_NOT_ALLOWED", "仅支持 PDF、DOCX、PPTX、TXT 和 Markdown")


def detect_kind(content: bytes, extension: str) -> str:
    if content.startswith(b"%PDF-"):
        return "pdf"
    if content.startswith(b"PK\x03\x04"):
        return inspect_ooxml(content)
    if content.startswith(b"\xd0\xcf\x11\xe0"):
        fail("FILE_TYPE_NOT_ALLOWED", "暂不支持旧版 DOC/PPT 或宏格式文档")
    try:
        content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ParseFailure(
            "FILE_TYPE_NOT_ALLOWED",
            "TXT 和 Markdown 文件必须使用 UTF-8 编码",
        ) from exc
    if b"\x00" in content:
        fail("FILE_TYPE_NOT_ALLOWED", "仅支持 PDF、DOCX、PPTX、TXT 和 Markdown")
    return "md" if extension == ".md" else "txt"


def pdf_page_tables(
    page, page_number: int
) -> tuple[str, list[ExtractedTable], list[Section], bool]:
    """Accept only a small, stable coordinate grid; never infer a missing cell."""
    fragments: list[tuple[float, float, str]] = []
    too_many_fragments = False
    unsupported_layout = False

    def visit(text, cm, tm, _font, _size):
        nonlocal too_many_fragments, unsupported_layout
        value = normalize_text(text)
        if not value:
            return
        if len(fragments) >= 10_000:
            too_many_fragments = True
            return
        if any(abs(matrix[1]) > 0.01 or abs(matrix[2]) > 0.01 for matrix in (cm, tm)):
            unsupported_layout = True
            return
        x, y = cm[4] + tm[4], cm[5] + tm[5]
        if math.isfinite(x) and math.isfinite(y):
            fragments.append((x, y, value))

    raw_text = normalize_text(page.extract_text(visitor_text=visit) or "")
    if too_many_fragments or unsupported_layout:
        # A rotated span or a crowded page invalidates the coordinate model,
        # but is not evidence that the entire page's prose should be discarded.
        return raw_text, [], [], True
    lines: list[list[tuple[float, float, str]]] = []
    for fragment in sorted(fragments, key=lambda item: (-item[1], item[0])):
        if lines and abs(lines[-1][0][1] - fragment[1]) <= 3:
            lines[-1].append(fragment)
        else:
            lines.append([fragment])
    candidates = [line for line in lines if len(line) >= 2]
    blocks: list[list[list[tuple[float, float, str]]]] = []
    for line in candidates:
        if blocks and 8 <= blocks[-1][-1][0][1] - line[0][1] <= 36:
            blocks[-1].append(line)
        else:
            blocks.append([line])
    tables: list[ExtractedTable] = []
    used: set[tuple[float, float, str]] = set()
    rejected: set[tuple[float, float, str]] = set()
    ambiguous = False
    for block in blocks:
        if len(block) < 2:
            continue
        ordered = [sorted(row, key=lambda item: item[0]) for row in block]
        # PDF text visitors also split ordinary sentences and contents entries
        # into positioned fragments. Require two complete, repeated column
        # grids before treating any part of a block as a potential table.
        anchor_pair = any(
            3 <= len(first) <= 8
            and len(first) == len(second)
            and all(first[i + 1][0] - first[i][0] >= 25 for i in range(len(first) - 1))
            and all(abs(a[0] - b[0]) <= 4 for a, b in zip(first, second, strict=True))
            for index, first in enumerate(ordered)
            for second in ordered[index + 1 : index + 3]
        )
        if not anchor_pair:
            continue
        columns = len(ordered[0])
        anchors = [item[0] for item in ordered[0]]
        valid = (
            3 <= columns <= 8
            and all(len(row) == columns for row in ordered)
            and all(anchors[i + 1] - anchors[i] >= 25 for i in range(columns - 1))
            and all(
                abs(row[column][0] - anchors[column]) <= 4
                for row in ordered
                for column in range(columns)
            )
        )
        if not valid:
            ambiguous = True
            rejected.update(fragment for row in ordered for fragment in row)
            continue
        locator = f"第 {page_number} 页 表格 {len(tables) + 1}"
        cells = tuple(
            TableCell(row_index, column_index, fragment[2])
            for row_index, row in enumerate(ordered, start=1)
            for column_index, fragment in enumerate(row, start=1)
        )
        tables.append(ExtractedTable(locator, len(ordered), columns, False, cells))
        used.update(fragment for row in ordered for fragment in row)
    remainder = [
        fragment[2] for fragment in fragments if fragment not in used and fragment not in rejected
    ]
    sections = (
        [Section(f"第 {page_number} 页", "\n".join(remainder))]
        if tables and remainder
        else []
    )
    return ("\n".join(remainder) if ambiguous else raw_text), tables, sections, ambiguous


def find_tesseract() -> str | None:
    executable = shutil.which("tesseract")
    if executable:
        return executable
    if sys.platform == "win32":
        standard = (
            Path(os.environ.get("ProgramFiles", r"C:\Program Files"))
            / "Tesseract-OCR"
            / "tesseract.exe"
        )
        if standard.is_file():
            return str(standard)
    return None


def tesseract_env() -> dict[str, str]:
    environment = os.environ.copy()
    local_data = Path(__file__).resolve().parents[2] / "data" / "tessdata"
    if (
        (local_data / "chi_sim.traineddata").is_file()
        and (local_data / "eng.traineddata").is_file()
    ):
        environment["TESSDATA_PREFIX"] = str(local_data)
    return environment


def local_tesseract() -> str | None:
    executable = find_tesseract()
    if not executable:
        return None
    kwargs = {"creationflags": subprocess.CREATE_NO_WINDOW} if sys.platform == "win32" else {}
    try:
        result = subprocess.run(
            [executable, "--list-langs"],
            capture_output=True,
            env=tesseract_env(),
            timeout=3,
            check=False,
            **kwargs,
        )
    except (OSError, subprocess.TimeoutExpired):
        return None
    languages = (result.stdout + result.stderr).decode("utf-8", errors="replace").splitlines()
    return executable if result.returncode == 0 and {"chi_sim", "eng"} <= set(languages) else None


def ocr_embedded_image(content: bytes, executable: str) -> str:
    if len(content) > MAX_OCR_IMAGE_BYTES:
        return ""
    try:
        with Image.open(io.BytesIO(content)) as image:
            if image.width * image.height > MAX_OCR_IMAGE_PIXELS or min(image.size) < 64:
                return ""
            output = io.BytesIO()
            image.convert("RGB").save(output, format="PNG")
    except (OSError, ValueError, UnidentifiedImageError, Image.DecompressionBombError):
        return ""
    kwargs = {"creationflags": subprocess.CREATE_NO_WINDOW} if sys.platform == "win32" else {}
    try:
        result = subprocess.run(
            [executable, "stdin", "stdout", "-l", "chi_sim+eng"],
            input=output.getvalue(),
            stdout=subprocess.PIPE,
            stderr=subprocess.DEVNULL,
            env=tesseract_env(),
            timeout=OCR_IMAGE_TIMEOUT_SECONDS,
            check=False,
            **kwargs,
        )
    except (OSError, subprocess.TimeoutExpired):
        return ""
    if result.returncode != 0:
        return ""
    return normalize_text(result.stdout.decode("utf-8", errors="replace"))[:MAX_OCR_TEXT_CHARS]


def parse_pdf(content: bytes, max_pages: int) -> tuple[list[Section], list[ExtractedTable], int]:
    try:
        reader = PdfReader(io.BytesIO(content), strict=True)
        if reader.is_encrypted:
            fail("FILE_ENCRYPTED", "暂不支持加密 PDF")
        page_count = len(reader.pages)
        if page_count > max_pages:
            fail("FILE_PAGE_LIMIT", f"PDF 最多支持 {max_pages} 页")
        sections = []
        tables = []
        executable: str | None = None
        ocr_checked = False
        ocr_count = 0
        saw_images = False
        has_readable_content = False
        for index, page in enumerate(reader.pages, start=1):
            text, page_tables, residual, ambiguous = pdf_page_tables(page, index)
            if page_tables:
                sections.extend(residual)
                for table in page_tables:
                    sections.extend(table_sections(table))
                tables.extend(page_tables)
                has_readable_content = True
            elif text:
                sections.append(Section(f"第 {index} 页", text))
                has_readable_content = True
            if ambiguous:
                sections.append(
                    Section(
                        f"第 {index} 页表格提示",
                        f"第 {index} 页：表格布局无法可靠识别，原始文字不能用于确认项目、"
                        "数值、单位的对应关系，请核对原件。",
                    )
                )
            image_count = len(page.images)
            if not image_count:
                continue
            saw_images = True
            if not ocr_checked:
                executable = local_tesseract()
                ocr_checked = True
            if executable is None:
                sections.append(
                    Section(
                        f"第 {index} 页图片提示",
                        f"第 {index} 页图片文字未读取：本地 Tesseract 或简体中文语言包不可用。",
                    )
                )
                continue
            processed = 0
            unreadable = 0
            for image_number, image in enumerate(page.images, start=1):
                if processed >= MAX_OCR_IMAGES_PER_PAGE or ocr_count >= MAX_OCR_IMAGES:
                    break
                ocr_count += 1
                processed += 1
                recognized = ocr_embedded_image(image.data, executable)
                if not recognized:
                    unreadable += 1
                    continue
                has_readable_content = True
                sections.append(
                    Section(
                        f"第 {index} 页图片 {image_number} OCR",
                        "图片 OCR 自动转录（未经人工核对；表格行列及单位对应未验证）：\n"
                        + recognized,
                    )
                )
            if image_count > processed:
                sections.append(
                    Section(
                        f"第 {index} 页图片提示",
                        f"第 {index} 页有 {image_count - processed} 张图片超出本地 OCR 处理上限，"
                        "其文字未读取。",
                    )
                )
            if unreadable:
                sections.append(
                    Section(
                        f"第 {index} 页图片提示",
                        f"第 {index} 页有 {unreadable} 张图片未识别出文字；"
                        "可能是非文字图片、内容不清晰或 OCR 超时。",
                    )
                )
        if not has_readable_content and saw_images and executable is None:
            fail("FILE_OCR_UNAVAILABLE", "本地 OCR 不可用，请安装 Tesseract 及简体中文语言包")
        if not has_readable_content:
            fail("FILE_NO_TEXT", "文档没有可读取文字，图片 OCR 也未识别到内容")
        return sections, tables, page_count
    except ParseFailure:
        raise
    except Exception as exc:
        raise ParseFailure("FILE_PARSE_FAILED", "PDF 已损坏或无法解析") from exc


def parse_docx(content: bytes) -> tuple[list[Section], list[ExtractedTable], None]:
    try:
        document = Document(io.BytesIO(content))
        sections: list[Section] = []
        tables: list[ExtractedTable] = []
        position = 1
        for paragraph in document.paragraphs:
            text = normalize_text(paragraph.text)
            if text:
                sections.append(Section(f"第 {position} 段", text))
                position += 1
        for table_index, table in enumerate(document.tables, start=1):
            cells = tuple(
                TableCell(row_index, column_index, normalize_text(cell.text))
                for row_index, row in enumerate(table.rows, start=1)
                for column_index, cell in enumerate(row.cells, start=1)
            )
            merged = len({cell._tc for row in table.rows for cell in row.cells}) != (
                len(table.rows) * len(table.columns)
            )
            extracted = ExtractedTable(
                f"表格 {table_index}", len(table.rows), len(table.columns), merged, cells
            )
            tables.append(extracted)
            sections.extend(table_sections(extracted))
        return sections, tables, None
    except Exception as exc:
        raise ParseFailure("FILE_PARSE_FAILED", "DOCX 已损坏或无法解析") from exc


def parse_pptx(content: bytes, max_pages: int) -> tuple[list[Section], list[ExtractedTable], int]:
    try:
        presentation = Presentation(io.BytesIO(content))
        page_count = len(presentation.slides)
        if page_count > max_pages:
            fail("FILE_PAGE_LIMIT", f"PPTX 最多支持 {max_pages} 张幻灯片")
        sections = []
        tables: list[ExtractedTable] = []
        for index, slide in enumerate(presentation.slides, start=1):
            parts = []
            for shape_index, shape in enumerate(slide.shapes, start=1):
                if getattr(shape, "has_text_frame", False):
                    text = normalize_text(shape.text)
                    if text:
                        parts.append(text)
                if getattr(shape, "has_table", False):
                    grid = shape.table
                    cells = tuple(
                        TableCell(row_index, column_index, normalize_text(cell.text))
                        for row_index, row in enumerate(grid.rows, start=1)
                        for column_index, cell in enumerate(row.cells, start=1)
                    )
                    merged = any(
                        cell.is_merge_origin or cell.is_spanned
                        for row in grid.rows
                        for cell in row.cells
                    )
                    extracted = ExtractedTable(
                        f"第 {index} 张幻灯片 表格 {shape_index}",
                        len(grid.rows),
                        len(grid.columns),
                        merged,
                        cells,
                    )
                    tables.append(extracted)
                    sections.extend(table_sections(extracted))
            if parts:
                sections.append(Section(f"第 {index} 张幻灯片", "\n".join(parts)))
        return sections, tables, page_count
    except ParseFailure:
        raise
    except Exception as exc:
        raise ParseFailure("FILE_PARSE_FAILED", "PPTX 已损坏或无法解析") from exc


def parse_text(content: bytes) -> tuple[list[Section], list[ExtractedTable], None]:
    text = normalize_text(content.decode("utf-8-sig"))
    blocks = [normalize_text(block) for block in re.split(r"\n\s*\n", text)]
    sections = [
        Section(f"第 {index} 段", block)
        for index, block in enumerate((item for item in blocks if item), start=1)
    ]
    return sections, [], None


def parse(content: bytes, extension: str, max_pages: int, max_chars: int) -> dict[str, object]:
    kind = detect_kind(content, extension)
    allowed_extensions = {
        "pdf": {".pdf"},
        "docx": {".docx"},
        "pptx": {".pptx"},
        "txt": {".txt"},
        "md": {".md"},
    }
    if extension not in allowed_extensions[kind]:
        fail("FILE_TYPE_NOT_ALLOWED", "文件扩展名与实际内容类型不一致")
    if kind == "pdf":
        sections, tables, page_count = parse_pdf(content, max_pages)
    elif kind == "docx":
        sections, tables, page_count = parse_docx(content)
    elif kind == "pptx":
        sections, tables, page_count = parse_pptx(content, max_pages)
    else:
        sections, tables, page_count = parse_text(content)
    char_count = sum(len(section.text) for section in sections)
    if char_count == 0:
        fail("FILE_NO_TEXT", "文档像是扫描件或没有可读取文字，暂时无法解析")
    if char_count > max_chars:
        fail("FILE_TEXT_TOO_LARGE", "文档文字过多，请拆分原文件后重试")
    return {
        "ok": True,
        "kind": kind,
        "pageCount": page_count,
        "charCount": char_count,
        "sections": [asdict(section) for section in sections],
        "tables": [asdict(table) for table in tables],
    }


def main() -> None:
    parser = argparse.ArgumentParser(add_help=False)
    parser.add_argument("--extension", required=True)
    parser.add_argument("--max-pages", type=int, required=True)
    parser.add_argument("--max-chars", type=int, required=True)
    parser.add_argument("--max-bytes", type=int, required=True)
    args = parser.parse_args()
    content = sys.stdin.buffer.read(args.max_bytes + 1)
    try:
        if not content:
            fail("FILE_PARSE_FAILED", "文档内容为空")
        if len(content) > args.max_bytes:
            fail("FILE_TOO_LARGE", f"单个文档不能超过 {args.max_bytes / (1024 * 1024):g} MB")
        result = parse(content, args.extension.lower(), args.max_pages, args.max_chars)
    except ParseFailure as exc:
        result = {"ok": False, "code": exc.code, "message": exc.message}
    except Exception:
        result = {"ok": False, "code": "FILE_PARSE_FAILED", "message": "文档无法解析"}
    # Keep the worker protocol ASCII-only so Windows console code pages cannot
    # corrupt JSON exchanged over the subprocess pipe.
    sys.stdout.write(json.dumps(result, ensure_ascii=True, separators=(",", ":")))


if __name__ == "__main__":
    main()
