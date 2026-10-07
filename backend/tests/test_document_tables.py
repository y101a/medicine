"""Table extraction keeps grid coordinates without guessing clinical meaning."""

import io

import pytest
from docx import Document
from fastapi import UploadFile
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.services.document_extract import extract_document


async def extract(name: str, content: bytes):
    return await extract_document(
        UploadFile(file=io.BytesIO(content), filename=name),
        max_bytes=10 * 1024 * 1024,
        max_pages=30,
        max_chars=500_000,
        max_chunk_chars=2_000,
        timeout_seconds=30,
    )


def docx_table_bytes(*, merged: bool = False) -> bytes:
    document = Document()
    grid = document.add_table(rows=2, cols=4)
    for column, text in enumerate(("项目", "结果", "单位", "参考范围")):
        grid.cell(0, column).text = text
    for column, text in enumerate(("白细胞", "6.2", "10^9/L", "3.5-9.5")):
        grid.cell(1, column).text = text
    if merged:
        grid.cell(0, 0).merge(grid.cell(0, 1))
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def pptx_table_bytes(*, merged: bool = False) -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    grid = slide.shapes.add_table(2, 4, Inches(1), Inches(1), Inches(8), Inches(1)).table
    for column, text in enumerate(("项目", "结果", "单位", "参考范围")):
        grid.cell(0, column).text = text
    for column, text in enumerate(("白细胞", "6.2", "10^9/L", "3.5-9.5")):
        grid.cell(1, column).text = text
    if merged:
        grid.cell(0, 0).merge(grid.cell(0, 1))
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def pdf_table_bytes(
    *, shifted: bool = False, missing: bool = False, missing_two: bool = False, pages: int = 1
) -> bytes:
    writer = PdfWriter()
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_ref = writer._add_object(font)
    for _page_number in range(pages):
        page = writer.add_blank_page(width=612, height=792)
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})}
        )
        commands = [b"BT /F1 12 Tf 72 760 Td (Lab report) Tj ET"]
        rows = (
            ("Item", "Result", "Unit", "Range"),
            ("WBC", "6.2", "10^9/L", "3.5-9.5"),
            ("RBC", "4.5", "10^12/L", "4.0-5.5"),
        )
        for row_index, row in enumerate(rows):
            for column_index, value in reversed(list(enumerate(row))):
                if missing and row_index == 1 and column_index == 1:
                    continue
                if missing_two and row_index == 1 and column_index in (1, 2):
                    continue
                x = (72, 210, 330, 450)[column_index]
                if shifted and row_index == 1 and column_index == 1:
                    x += 18
                y = 700 - row_index * 24
                commands.append(f"BT /F1 12 Tf {x} {y} Td ({value}) Tj ET".encode())
        stream = DecodedStreamObject()
        stream.set_data(b"\n".join(commands))
        page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def pdf_fragmented_prose_bytes(*, contents: bool = False) -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    rows = (
        ((72, "Experiment"), (155, "one"), (210, "setup"), (450, "3")),
        ((72, "Experiment"), (155, "two"), (210, "results"), (450, "6")),
        ((72, "Experiment"), (155, "three"), (210, "summary"), (450, "9")),
    ) if contents else (
        ((72, "This"), (98, "is"), (116, "a"), (131, "paragraph")),
        ((72, "The"), (96, "layout"), (139, "contains"), (191, "separate"), (247, "words")),
        ((72, "It"), (86, "is"), (103, "not"), (127, "a"), (142, "table")),
    )
    commands = []
    for row_index, row in enumerate(rows):
        y = 700 - row_index * (60 if contents else 20)
        for x, value in row:
            commands.append(f"BT /F1 12 Tf {x} {y} Td ({value}) Tj ET".encode())
    stream = DecodedStreamObject()
    stream.set_data(b"\n".join(commands))
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


@pytest.mark.parametrize(
    ("name", "factory"),
    [("report.docx", docx_table_bytes), ("report.pptx", pptx_table_bytes)],
)
async def test_table_cells_keep_row_and_column(name, factory):
    result = await extract(name, factory())
    assert len(result.tables) == 1
    table = result.tables[0]
    assert (table.row_count, table.column_count, table.layout_uncertain) == (2, 4, False)
    assert [(cell.row, cell.column, cell.text) for cell in table.cells[4:]] == [
        (2, 1, "白细胞"),
        (2, 2, "6.2"),
        (2, 3, "10^9/L"),
        (2, 4, "3.5-9.5"),
    ]
    row_chunk = next(chunk for chunk in result.chunks if "白细胞" in chunk.text)
    assert "第 2 行" in row_chunk.locator
    assert "第 1 列：白细胞" in row_chunk.text
    assert "第 2 列：6.2" in row_chunk.text
    assert "第 3 列：10^9/L" in row_chunk.text
    assert "第 4 列：3.5-9.5" in row_chunk.text


@pytest.mark.parametrize(
    ("name", "factory"),
    [("report.docx", docx_table_bytes), ("report.pptx", pptx_table_bytes)],
)
async def test_merged_table_marks_layout_uncertain(name, factory):
    result = await extract(name, factory(merged=True))
    assert result.tables[0].layout_uncertain is True
    assert any("行列对应需核对" in chunk.text for chunk in result.chunks)


async def test_pdf_table_uses_positions_not_text_stream_order():
    result = await extract("report.pdf", pdf_table_bytes())
    assert len(result.tables) == 1
    table = result.tables[0]
    assert (table.row_count, table.column_count, table.layout_uncertain) == (3, 4, False)
    assert [(cell.row, cell.column, cell.text) for cell in table.cells[4:8]] == [
        (2, 1, "WBC"),
        (2, 2, "6.2"),
        (2, 3, "10^9/L"),
        (2, 4, "3.5-9.5"),
    ]
    assert any("第 1 页 表格 1 第 2 行" in chunk.locator for chunk in result.chunks)


async def test_pdf_misaligned_table_does_not_guess_columns():
    result = await extract("report.pdf", pdf_table_bytes(shifted=True))
    assert result.tables == ()
    assert any("表格布局无法可靠识别" in chunk.text for chunk in result.chunks)
    assert all("6.2" not in chunk.text for chunk in result.chunks)


async def test_pdf_missing_cell_does_not_join_surrounding_rows():
    result = await extract("report.pdf", pdf_table_bytes(missing=True))
    assert result.tables == ()
    assert any("表格布局无法可靠识别" in chunk.text for chunk in result.chunks)
    assert all("6.2" not in chunk.text for chunk in result.chunks)


async def test_pdf_short_row_does_not_join_surrounding_rows():
    result = await extract("report.pdf", pdf_table_bytes(missing_two=True))
    assert result.tables == ()
    assert any("表格布局无法可靠识别" in chunk.text for chunk in result.chunks)


async def test_pdf_tables_do_not_merge_across_pages():
    result = await extract("report.pdf", pdf_table_bytes(pages=2))
    assert len(result.tables) == 2
    assert [table.locator for table in result.tables] == ["第 1 页 表格 1", "第 2 页 表格 1"]


@pytest.mark.parametrize("contents", [False, True], ids=["prose", "contents"])
async def test_pdf_fragmented_non_table_text_is_preserved(contents):
    result = await extract("manual.pdf", pdf_fragmented_prose_bytes(contents=contents))
    assert result.tables == ()
    assert not any("表格布局无法可靠识别" in chunk.text for chunk in result.chunks)
    joined = " ".join(chunk.text for chunk in result.chunks)
    assert ("Experiment" if contents else "paragraph") in joined
    assert ("summary" if contents else "separate") in joined
