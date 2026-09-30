import importlib
import os
from datetime import UTC, datetime, timedelta
from io import BytesIO
from uuid import UUID, uuid4
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches
from sqlalchemy import select

from app.core.config import get_settings
from app.core.errors import AppError
from app.core.security import create_access_token
from app.models import (
    GeneratedDocument,
    GeneratedDocumentDeletionJob,
    GeneratedDocumentFile,
    GeneratedDocumentTemplate,
)
from app.rag.retrieval import RetrievalCitation, RetrievalResult
from app.services import document_generations as generation_service
from app.services.auth import ACCESS_COOKIE
from app.services.document_generation_content import DocumentContentPlan
from app.services.document_generation_intent import medical_document_generation_request
from app.services.document_renderers import render_docx, render_pptx
from app.services.generated_document_deletion import (
    process_document_deletion_jobs,
    reconcile_orphaned_generated_documents,
)
from app.services.generated_document_validation import validate_generated_document
from app.services.presentation_template import validate_presentation_template


def _plan() -> DocumentContentPlan:
    return DocumentContentPlan.model_validate(
        {
            "title": "高血压预防科普",
            "audience": "普通成年人",
            "summary": "高血压预防需要长期关注生活方式、规范测量血压并按需接受医疗评估。",
            "sections": [
                {
                    "heading": "认识高血压",
                    "paragraphs": ["血压水平需要结合规范测量结果进行判断。"],
                    "bullets": [],
                    "sourceIds": [1],
                },
                {
                    "heading": "日常预防",
                    "paragraphs": [],
                    "bullets": ["保持规律运动和合理饮食。", "按建议记录血压变化。"],
                    "sourceIds": [1],
                },
            ],
            "warningSigns": [{"text": "出现胸痛、呼吸困难等表现时应及时就医。", "sourceIds": [1]}],
        }
    )


def _citation() -> RetrievalCitation:
    return RetrievalCitation(
        chunk_id=uuid4(),
        document_id=uuid4(),
        source_key="developer_prototype_knowledge",
        organization="项目开发者",
        title="高血压原型资料",
        canonical_url="https://prototype.invalid/medical-knowledge/hypertension",
        section_title="预防",
        version_label="test",
        published_at=None,
        content="高血压预防包括合理饮食、规律运动和规范测量。",
        similarity=0.9,
        trust_level="prototype",
    )


@pytest.mark.parametrize(
    ("content", "formats"),
    [
        ("帮我生成一份高血压预防 PPT", ("pptx",)),
        ("制作胃食管反流 Word 科普文档", ("docx",)),
        ("把哮喘预防做成 PPT 和 Word", ("docx", "pptx")),
    ],
)
def test_document_generation_intent_requires_explicit_format(content, formats):
    result = medical_document_generation_request(content)
    assert result is not None
    assert result.formats == formats


@pytest.mark.parametrize(
    "content",
    [
        "你能生成 PPT 吗？",
        "不要生成高血压 Word 文档",
        "生成一份年度销售 PPT",
        "生成一份人体结构科普资料",
        "介绍一下人体结构",
    ],
)
def test_non_document_generation_intent(content):
    assert medical_document_generation_request(content) is None


@pytest.mark.parametrize(
    ("content", "expected"),
    [
        ("帮我生成一份高血压预防 PPT", False),
        ("帮我使用模板生成一份高血压预防 PPT", True),
        ("帮我用内置模板生成高血压预防 PPT", True),
        ("帮我生成高血压预防 PPT，不需要模板", False),
        ("帮我生成高血压预防 Word 文档，使用模板", False),
    ],
)
def test_document_generation_template_intent(content, expected):
    request = medical_document_generation_request(content)
    assert request is not None
    assert request.template_requested is expected
    assert request.topic == "高血压预防"


def test_renderers_create_safe_reopenable_office_files():
    sources = [
        {
            "referenceNumber": 1,
            "organization": "项目开发者",
            "title": "高血压原型资料",
            "canonicalUrl": "https://prototype.invalid/medical-knowledge/hypertension",
        }
    ]
    docx = render_docx(_plan(), sources)
    pptx, slides = render_pptx(_plan(), sources)
    validate_generated_document(docx, "docx", max_bytes=10 * 1024 * 1024)
    validate_generated_document(pptx, "pptx", max_bytes=10 * 1024 * 1024)
    assert Document(BytesIO(docx)).paragraphs[0].text == "高血压预防科普"
    assert len(Presentation(BytesIO(pptx)).slides) == slides
    assert slides >= 5


def test_builtin_ppt_template_changes_background_without_changing_default():
    default_bytes, _ = render_pptx(_plan(), [])
    styled_bytes, _ = render_pptx(_plan(), [], template_mode="builtin")
    default_deck = Presentation(BytesIO(default_bytes))
    styled_deck = Presentation(BytesIO(styled_bytes))

    assert default_deck.slides[0].background.fill.type != styled_deck.slides[0].background.fill.type
    assert styled_deck.slides[0].background.fill.fore_color.rgb is not None
    assert styled_deck.slides[1].background.fill.fore_color.rgb is not None
    assert len(styled_deck.slides) == len(default_deck.slides)
    validate_generated_document(styled_bytes, "pptx", max_bytes=10 * 1024 * 1024)


def _user_template() -> bytes:
    deck = Presentation()
    deck.slide_width = Inches(12)
    deck.slide_height = Inches(7)
    picture = Image.new("RGB", (100, 40), (18, 80, 140))
    image = BytesIO()
    picture.save(image, format="PNG")
    for index in range(2):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(239, 238, 220)
        slide.shapes.add_picture(
            BytesIO(image.getvalue()), Inches(0), Inches(0), width=Inches(12), height=Inches(0.8)
        )
        title = slide.shapes.add_textbox(Inches(0.6), Inches(1), Inches(10), Inches(0.9))
        title.text = f"旧模板私有文字{index}"
        body = slide.shapes.add_textbox(Inches(0.6), Inches(2), Inches(10), Inches(4))
        body.text = "旧模板正文"
    output = BytesIO()
    deck.save(output)
    return output.getvalue()


def test_uploaded_ppt_template_keeps_background_picture_and_layout():
    template = _user_template()
    validate_presentation_template(template)
    output, count = render_pptx(_plan(), [], template_bytes=template)
    deck = Presentation(BytesIO(output))
    assert len(deck.slides) == count
    assert deck.slide_width == Inches(12)
    assert deck.slides[0].background.fill.fore_color.rgb == RGBColor(239, 238, 220)
    assert any(shape.shape_type == 13 for shape in deck.slides[0].shapes)
    assert any(
        shape.left == Inches(0.6) and shape.top == Inches(1) for shape in deck.slides[1].shapes
    )
    all_text = " ".join(
        shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame
    )
    assert "高血压预防科普" in all_text
    assert "旧模板私有文字" not in all_text
    assert "不作诊断依据" in all_text
    validate_generated_document(output, "pptx", max_bytes=10 * 1024 * 1024)


def test_uploaded_ppt_template_requires_cover_and_body():
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    payload = BytesIO()
    deck.save(payload)
    with pytest.raises(AppError) as exc:
        validate_presentation_template(payload.getvalue())
    assert exc.value.code == "PPT_TEMPLATE_INVALID"


def test_uploaded_ppt_template_rejects_active_parts_and_oversize():
    with pytest.raises(AppError) as too_large:
        validate_presentation_template(b"x" * (10 * 1024 * 1024 + 1))
    assert too_large.value.code == "PPT_TEMPLATE_TOO_LARGE"

    modified = BytesIO()
    with (
        ZipFile(BytesIO(_user_template())) as source,
        ZipFile(modified, "w", ZIP_DEFLATED) as target,
    ):
        for part in source.infolist():
            target.writestr(part, source.read(part))
        target.writestr("ppt/vbaProject.bin", b"macro")
    with pytest.raises(AppError) as forbidden:
        validate_presentation_template(modified.getvalue())
    assert forbidden.value.code == "PPT_TEMPLATE_INVALID"


async def _create_request(
    client,
    user_id,
    content="帮我生成一份高血压预防 PPT 和 Word",
):
    client.cookies.set(ACCESS_COOKIE, create_access_token(user_id))
    consultation = await client.post(
        "/api/v1/consultations",
        json={"clientRequestId": str(uuid4())},
    )
    cid = UUID(consultation.json()["id"])
    message = await client.post(
        f"/api/v1/consultations/{cid}/messages",
        json={
            "clientRequestId": str(uuid4()),
            "content": content,
        },
    )
    assert message.status_code == 201
    assert message.json()["documentGenerationIntent"] is True
    assert message.json()["imageGenerationIntent"] is False
    return cid, UUID(message.json()["id"])


async def test_create_is_idempotent_notice_guarded_and_user_isolated(profile_context):
    client, db, (user, other) = profile_context
    cid, message_id = await _create_request(client, user.id)
    payload = {
        "clientRequestId": str(uuid4()),
        "sourceMessageId": str(message_id),
        "noticeAccepted": True,
        "noticeVersion": "document-generation-v1",
    }
    missing = await client.post(
        f"/api/v1/consultations/{cid}/document-generations",
        json={**payload, "noticeAccepted": False},
    )
    assert missing.status_code == 422
    assert missing.json()["code"] == "DOCUMENT_GENERATION_NOTICE_REQUIRED"
    first = await client.post(
        f"/api/v1/consultations/{cid}/document-generations",
        json=payload,
    )
    second = await client.post(
        f"/api/v1/consultations/{cid}/document-generations",
        json=payload,
    )
    replayed_source = await client.post(
        f"/api/v1/consultations/{cid}/document-generations",
        json={**payload, "clientRequestId": str(uuid4())},
    )
    assert first.status_code == second.status_code == 202
    assert first.json()["id"] == second.json()["id"]
    assert replayed_source.status_code == 202
    assert replayed_source.json()["id"] == first.json()["id"]
    assert first.json()["formats"] == ["docx", "pptx"]
    row = await db.get(GeneratedDocument, UUID(first.json()["id"]))
    assert row is not None and row.status == "queued"
    assert row.title_ciphertext is None

    client.cookies.set(ACCESS_COOKIE, create_access_token(other.id))
    hidden = await client.get(f"/api/v1/document-generations/{row.id}")
    assert hidden.status_code == 404


async def test_uploaded_template_is_encrypted_used_and_deleted(
    profile_context, tmp_path, monkeypatch
):
    client, db, (user, _) = profile_context
    monkeypatch.setattr(get_settings(), "document_generation_storage_dir", str(tmp_path))
    cid, message_id = await _create_request(client, user.id, "帮我生成一份高血压预防 PPT")
    created = await client.post(
        f"/api/v1/consultations/{cid}/document-generations/with-template",
        data={
            "clientRequestId": str(uuid4()),
            "sourceMessageId": str(message_id),
            "noticeAccepted": "true",
            "noticeVersion": "document-generation-v1",
            "templateNoticeAccepted": "true",
        },
        files={
            "file": (
                "mine.pptx",
                _user_template(),
                "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            )
        },
    )
    assert created.status_code == 202, created.text
    assert created.json()["templateMode"] == "uploaded"
    missing_template_notice = await client.post(
        f"/api/v1/consultations/{cid}/document-generations/with-template",
        data={
            "clientRequestId": str(uuid4()),
            "sourceMessageId": str(message_id),
            "noticeAccepted": "true",
            "noticeVersion": "document-generation-v1",
            "templateNoticeAccepted": "false",
        },
        files={"file": ("mine.pptx", _user_template(), "application/octet-stream")},
    )
    assert missing_template_notice.status_code == 422
    assert missing_template_notice.json()["code"] == "PPT_TEMPLATE_NOTICE_REQUIRED"
    different_template = await client.post(
        f"/api/v1/consultations/{cid}/document-generations/with-template",
        data={
            "clientRequestId": str(uuid4()),
            "sourceMessageId": str(message_id),
            "noticeAccepted": "true",
            "noticeVersion": "document-generation-v1",
            "templateNoticeAccepted": "true",
        },
        files={"file": ("different.pptx", b"different", "application/octet-stream")},
    )
    assert different_template.status_code == 409
    assert different_template.json()["code"] == "DOCUMENT_GENERATION_TEMPLATE_CONFLICT"
    generation_id = UUID(created.json()["id"])
    template = await db.get(GeneratedDocumentTemplate, generation_id)
    assert template is not None
    encrypted_path = tmp_path / template.storage_key
    assert encrypted_path.exists() and encrypted_path.read_bytes()[:2] != b"PK"
    assert (
        await reconcile_orphaned_generated_documents(
            db, older_than=datetime.now(UTC) + timedelta(days=1)
        )
        == 0
    )

    async def fake_search(_db, query, **kwargs):
        return RetrievalResult(
            query_redacted=False, model="test", version="test", items=[_citation()]
        )

    async def fake_plan(topic, citations):
        return _plan()

    monkeypatch.setattr(generation_service, "search_knowledge", fake_search)
    monkeypatch.setattr(generation_service, "generate_content_plan", fake_plan)
    assert await generation_service.process_one(db, generation_id)
    result = await client.get(f"/api/v1/document-generations/{generation_id}")
    pptx = next(item for item in result.json()["files"] if item["format"] == "pptx")
    download = await client.get(
        f"/api/v1/document-generations/{generation_id}/files/{pptx['id']}/content"
    )
    assert download.status_code == 200
    deck = Presentation(BytesIO(download.content))
    assert deck.slide_width == Inches(12)
    assert deck.slides[0].background.fill.fore_color.rgb == RGBColor(239, 238, 220)
    removed = await client.delete(f"/api/v1/document-generations/{generation_id}")
    assert removed.status_code == 204
    assert not encrypted_path.exists()


@pytest.mark.parametrize(
    ("request_content", "uses_builtin_template"),
    [
        ("帮我生成一份高血压预防 PPT 和 Word", False),
        ("帮我用内置模板生成一份高血压预防 PPT 和 Word", True),
    ],
)
async def test_worker_generates_encrypted_files_downloads_and_deletes(
    profile_context,
    tmp_path,
    monkeypatch,
    request_content,
    uses_builtin_template,
):
    client, db, (user, _) = profile_context
    settings = get_settings()
    monkeypatch.setattr(settings, "document_generation_storage_dir", str(tmp_path))
    cid, message_id = await _create_request(client, user.id, request_content)
    created = await client.post(
        f"/api/v1/consultations/{cid}/document-generations",
        json={
            "clientRequestId": str(uuid4()),
            "sourceMessageId": str(message_id),
            "noticeAccepted": True,
            "noticeVersion": "document-generation-v1",
        },
    )
    generation_id = UUID(created.json()["id"])

    async def fake_search(_db, query, **kwargs):
        assert "高血压" in query
        return RetrievalResult(
            query_redacted=False,
            model="test",
            version="test",
            items=[_citation()],
        )

    async def fake_plan(topic, citations):
        assert "高血压" in topic
        assert len(citations) == 1
        return _plan()

    monkeypatch.setattr(generation_service, "search_knowledge", fake_search)
    monkeypatch.setattr(generation_service, "generate_content_plan", fake_plan)
    assert await generation_service.process_one(db, generation_id) is True

    status = await client.get(f"/api/v1/document-generations/{generation_id}")
    assert status.status_code == 200
    payload = status.json()
    assert payload["status"] == "completed"
    assert payload["title"] == "高血压预防科普"
    assert payload["sourceCount"] == 1
    assert payload["expiresAt"] is not None
    assert [item["format"] for item in payload["files"]] == ["docx", "pptx"]
    assert payload["files"][0]["pageCount"] is None
    assert payload["files"][1]["pageCount"] >= 5

    rows = (
        await db.scalars(
            select(GeneratedDocumentFile).where(
                GeneratedDocumentFile.generation_id == generation_id
            )
        )
    ).all()
    assert len(rows) == 2
    for row in rows:
        stored = tmp_path / row.storage_key
        assert stored.exists()
        assert stored.read_bytes()[:2] != b"PK"
        response = await client.get(
            f"/api/v1/document-generations/{generation_id}/files/{row.id}/content"
        )
        assert response.status_code == 200
        assert response.content[:2] == b"PK"
        if row.format == "pptx":
            deck = Presentation(BytesIO(response.content))
            background_type = deck.slides[0].background.fill.type
            assert (background_type.name == "SOLID") is uses_builtin_template
        assert response.headers["cache-control"] == "no-store, private"
        assert response.headers["content-disposition"].startswith("attachment;")

    listed = await client.get(f"/api/v1/consultations/{cid}/document-generations")
    assert listed.status_code == 200
    assert listed.json()["items"][0]["sourceMessageId"] == str(message_id)

    deleted = await client.delete(f"/api/v1/document-generations/{generation_id}")
    assert deleted.status_code == 204
    assert all(not (tmp_path / row.storage_key).exists() for row in rows)


async def test_worker_fails_without_evidence(profile_context, monkeypatch):
    client, db, (user, _) = profile_context
    cid, message_id = await _create_request(client, user.id)
    created = await client.post(
        f"/api/v1/consultations/{cid}/document-generations",
        json={
            "clientRequestId": str(uuid4()),
            "sourceMessageId": str(message_id),
            "noticeAccepted": True,
            "noticeVersion": "document-generation-v1",
        },
    )
    generation_id = UUID(created.json()["id"])

    async def no_results(*args, **kwargs):
        return RetrievalResult(
            query_redacted=False,
            model="test",
            version="test",
            items=[],
        )

    monkeypatch.setattr(generation_service, "search_knowledge", no_results)
    assert await generation_service.process_one(db, generation_id) is False
    row = await db.get(GeneratedDocument, generation_id)
    await db.refresh(row)
    assert row.status == "failed"
    assert row.error_code == "DOCUMENT_EVIDENCE_INSUFFICIENT"


async def test_worker_fails_when_all_retrieved_evidence_is_below_gate(
    profile_context, monkeypatch
):
    client, db, (user, _) = profile_context
    cid, message_id = await _create_request(client, user.id)
    created = await client.post(
        f"/api/v1/consultations/{cid}/document-generations",
        json={
            "clientRequestId": str(uuid4()),
            "sourceMessageId": str(message_id),
            "noticeAccepted": True,
            "noticeVersion": "document-generation-v1",
        },
    )
    generation_id = UUID(created.json()["id"])

    weak = _citation()
    weak = RetrievalCitation(
        chunk_id=weak.chunk_id,
        document_id=weak.document_id,
        source_key=weak.source_key,
        organization=weak.organization,
        title=weak.title,
        canonical_url=weak.canonical_url,
        section_title=weak.section_title,
        version_label=weak.version_label,
        published_at=weak.published_at,
        content=weak.content,
        similarity=0.55,
        trust_level=weak.trust_level,
        vector_similarity=0.55,
    )

    async def weak_results(*args, **kwargs):
        return RetrievalResult(
            query_redacted=False,
            model="test",
            version="test",
            items=[weak],
        )

    monkeypatch.setattr(generation_service, "search_knowledge", weak_results)
    assert await generation_service.process_one(db, generation_id) is False
    row = await db.get(GeneratedDocument, generation_id)
    await db.refresh(row)
    assert row.status == "failed"
    assert row.error_code == "DOCUMENT_EVIDENCE_INSUFFICIENT"


async def test_worker_does_not_reclaim_queued_task_with_live_lease(profile_context):
    client, db, (user, _) = profile_context
    cid, message_id = await _create_request(client, user.id)
    created = await client.post(
        f"/api/v1/consultations/{cid}/document-generations",
        json={
            "clientRequestId": str(uuid4()),
            "sourceMessageId": str(message_id),
            "noticeAccepted": True,
            "noticeVersion": "document-generation-v1",
        },
    )
    generation_id = UUID(created.json()["id"])
    row = await db.get(GeneratedDocument, generation_id)
    row.lease_owner = "another-worker"
    row.lease_expires_at = datetime.now(UTC) + timedelta(minutes=2)
    await db.commit()

    assert await generation_service.process_due(db, limit=10) == (0, 0)
    await db.refresh(row)
    assert row.status == "queued"
    assert row.lease_owner == "another-worker"


async def test_orphaned_generated_document_files_are_queued_and_deleted(
    profile_context,
    tmp_path,
    monkeypatch,
):
    _, db, _ = profile_context
    monkeypatch.setattr(get_settings(), "document_generation_storage_dir", str(tmp_path))
    storage_dir = tmp_path / "files" / "aa"
    storage_dir.mkdir(parents=True)
    orphan = storage_dir / "orphan.bin"
    temporary = storage_dir / "interrupted.tmp"
    orphan.write_bytes(b"encrypted")
    temporary.write_bytes(b"partial")
    stale = (datetime.now(UTC) - timedelta(hours=2)).timestamp()
    os.utime(orphan, (stale, stale))
    os.utime(temporary, (stale, stale))

    found = await reconcile_orphaned_generated_documents(
        db,
        older_than=datetime.now(UTC) - timedelta(hours=1),
    )
    assert found == 2
    jobs = (await db.scalars(select(GeneratedDocumentDeletionJob))).all()
    assert {job.storage_key for job in jobs} == {
        "files/aa/orphan.bin",
        "files/aa/interrupted.tmp",
    }

    deleted = await process_document_deletion_jobs(db)
    assert deleted.completed == 2
    assert not orphan.exists()
    assert not temporary.exists()


def test_document_generation_migration_refuses_unsafe_downgrade(monkeypatch):
    migration = importlib.import_module("migrations.versions.20260927_0024_generated_documents")

    class Bind:
        @staticmethod
        def scalar(statement):
            assert "generated_documents" in str(statement)
            return 1

    monkeypatch.setattr(migration.op, "get_bind", lambda: Bind())
    with pytest.raises(RuntimeError, match="generated document data exists"):
        migration.downgrade()


def test_generated_document_default_retention_is_30_days():
    assert get_settings().document_generation_retention_days == 30


def test_content_plan_rejects_unknown_source_number():
    invalid = _plan().model_copy(deep=True)
    invalid.sections[0].source_ids = [2]
    from app.services.document_generation_content import _validate_source_ids

    with pytest.raises(AppError) as raised:
        _validate_source_ids(invalid, 1)
    assert raised.value.code == "DOCUMENT_MODEL_INVALID_RESPONSE"
